"""G1 input: turn a presentation or research doc into films/<name>/source.md (plain Markdown, kept with the film).

.pptx → one section per slide: title, text frames, tables, speaker notes.  .pdf → text per page.  .md/.txt → copied.
"""

import datetime
from pathlib import Path


def pptx_md(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None else ""
        out.append(f"## Slide {i}" + (f" — {title}" if title else ""))
        for sh in slide.shapes:
            if sh == slide.shapes.title:
                continue
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        out.append("  " * para.level + f"- {t}")
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    out.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in row.cells) + " |")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append("\n> Notes: " + notes.replace("\n", "\n> "))
        out.append("")
    return "\n".join(out)


def pdf_md(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return "\n".join(f"## Page {i}\n\n{(page.extract_text() or '').strip()}\n" for i, page in enumerate(reader.pages, 1))


def extract(src: Path, dest_dir: Path) -> Path:
    ext = src.suffix.lower()
    if ext == ".pptx":
        body = pptx_md(src)
    elif ext == ".pdf":
        body = pdf_md(src)
    elif ext in (".md", ".txt"):
        body = src.read_text(encoding="utf-8")
    else:
        raise SystemExit(f"Unsupported source {src.suffix} (pptx, pdf, md, txt)")
    dest_dir.mkdir(parents=True, exist_ok=True)
    out = dest_dir / "source.md"
    head = f"# Source — {src.name}\n\n*Extracted {datetime.date.today().isoformat()} by `studio extract` from `{src}`.*\n\n"
    if out.exists():  # several sources: append
        out.write_text(out.read_text(encoding="utf-8") + "\n\n---\n\n" + head + body, encoding="utf-8")
    else:
        out.write_text(head + body, encoding="utf-8")
    return out
